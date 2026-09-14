from fastapi import APIRouter, Depends, HTTPException, status, Query, Response, UploadFile, File
from app.db.mongodb import get_database
from app.schemas.product import ProductCreate, ProductUpdate, ProductResponse
from app.api.v1.auth import get_current_user
from bson import ObjectId
from datetime import datetime
from typing import Optional, List
import re
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

router = APIRouter()

def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[\s_-]+', '-', text)
    return text.strip('-')

def parse_denominazione_acronym(denom: str) -> str:
    if not denom:
        return "DOC"
    d_upper = str(denom).strip().upper()
    if "DOCG" in d_upper or "D.O.C.G." in d_upper:
        return "DOCG"
    elif "DOC" in d_upper or "D.O.C." in d_upper:
        return "DOC"
    elif "DOP" in d_upper or "D.O.P." in d_upper:
        return "DOP"
    elif "IGP" in d_upper or "I.G.P." in d_upper:
        return "IGP"
    elif "IGT" in d_upper or "I.G.T." in d_upper:
        return "IGT"
    return "DOC"

def clean_wine_title(title: str) -> str:
    if not title:
        return ""
    # Remove denominations (DOCG, DOC, DOP, IGP, IGT with or without dots)
    cleaned = re.sub(
        r'(?i)\b(?:d[\s.]*o[\s.]*c[\s.]*g|d[\s.]*o[\s.]*c|d[\s.]*o[\s.]*p|i[\s.]*g[\s.]*p|i[\s.]*g[\s.]*t)\b\.?',
        '',
        title
    )
    # Remove 4-digit years (e.g. 19xx, 20xx)
    cleaned = re.sub(r'\b(19\d{2}|20\d{2})\b', '', cleaned)
    # Clean extra spaces and punctuation
    cleaned = re.sub(r'\s+', ' ', cleaned).strip(' -.,')
    return cleaned if cleaned else title.strip()

def clean_product_slug(text: str) -> str:
    if not text:
        return ""
    # Normalize hyphens and underscores to spaces so word boundary regex works
    cleaned = str(text).replace('-', ' ').replace('_', ' ')
    
    # Remove denominations
    cleaned = re.sub(
        r'(?i)\b(?:d[\s.]*o[\s.]*c[\s.]*g|d[\s.]*o[\s.]*c|d[\s.]*o[\s.]*p|i[\s.]*g[\s.]*p|i[\s.]*g[\s.]*t)\b\.?',
        '',
        cleaned
    )
    
    # Remove 4-digit years
    cleaned = re.sub(r'\b(19\d{2}|20\d{2})\b', '', cleaned)
    
    # Slugify in kebab-case
    cleaned = cleaned.lower().strip()
    cleaned = re.sub(r'[^\w\s-]', '', cleaned)
    cleaned = re.sub(r'[\s_-]+', '-', cleaned).strip('-')
    
    if not cleaned:
        fallback = str(text).lower().strip()
        fallback = re.sub(r'[^\w\s-]', '', fallback)
        cleaned = re.sub(r'[\s_-]+', '-', fallback).strip('-')
        
    return cleaned

async def generate_unique_product_slug(
    db, 
    name: str, 
    producer_id, 
    is_riserva: bool = False,
    exclude_id = None
) -> str:
    cleaned_name = clean_wine_title(name)
    ris_tag = "riserva" if is_riserva and "riserva" not in cleaned_name.lower() else ""
    base_slug = clean_product_slug(f"{cleaned_name} {ris_tag}")
    if not base_slug:
        base_slug = "vino"
    
    # 1. Check if base_slug is free
    query = {"slug": base_slug}
    if exclude_id:
        query["_id"] = {"$ne": ObjectId(exclude_id)}
        
    existing = await db.products.find_one(query)
    if not existing:
        return base_slug
        
    if exclude_id and str(existing.get("_id")) == str(exclude_id):
        return base_slug
        
    # Collision! Retrieve producer slug
    producer = None
    if producer_id and ObjectId.is_valid(producer_id):
        producer = await db.producers.find_one({"_id": ObjectId(producer_id)})
    prod_slug = producer.get("slug") if producer else ""
    
    # 2. If existing product belongs to another winery, append this winery's slug
    if prod_slug and str(existing.get("producer_id")) != str(producer_id):
        candidate = f"{base_slug}-{prod_slug}"
        c_query = {"slug": candidate}
        if exclude_id:
            c_query["_id"] = {"$ne": ObjectId(exclude_id)}
        if not await db.products.find_one(c_query):
            return candidate

    # 3. Fallback to incremental counter (-2, -3, ...)
    counter = 2
    while True:
        candidate = f"{base_slug}-{counter}"
        c_query = {"slug": candidate}
        if exclude_id:
            c_query["_id"] = {"$ne": ObjectId(exclude_id)}
        if not await db.products.find_one(c_query):
            return candidate
        counter += 1

from html.parser import HTMLParser
import urllib.request

class SimpleHTMLScraper(HTMLParser):
    def __init__(self):
        super().__init__()
        self.title = ""
        self.og_title = ""
        self.og_desc = ""
        self.og_image = ""
        self.paragraphs = []
        self.in_title = False
        self.in_p = False
        self.current_p = ""

    def handle_starttag(self, tag, attrs):
        attr_dict = dict(attrs)
        if tag == "title":
            self.in_title = True
        elif tag == "meta":
            prop = attr_dict.get("property", "") or attr_dict.get("name", "")
            content = attr_dict.get("content", "")
            if prop == "og:title":
                self.og_title = content
            elif prop in ["og:description", "description"]:
                self.og_desc = content
            elif prop == "og:image":
                self.og_image = content
        elif tag == "p":
            self.in_p = True
            self.current_p = ""

    def handle_endtag(self, tag):
        if tag == "title":
            self.in_title = False
        elif tag == "p":
            self.in_p = False
            if len(self.current_p.strip()) > 30:
                self.paragraphs.append(self.current_p.strip())

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        if self.in_p:
            self.current_p += data

async def format_product_response(doc: dict, db) -> dict:
    doc["id"] = str(doc["_id"])
    doc["producer_id"] = str(doc["producer_id"])
    
    producer = await db.producers.find_one({"_id": ObjectId(doc["producer_id"])})
    if producer:
        doc["producer_name"] = producer.get("company_name", "")
        doc["producer_slug"] = producer.get("slug", "")
    return doc

@router.post("/scrape-url")
async def scrape_url(payload: dict):
    url = payload.get("url", "").strip()
    if not url or not (url.startswith("http://") or url.startswith("https://")):
        raise HTTPException(status_code=400, detail="URL non valido")
    
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
        with urllib.request.urlopen(req, timeout=10) as response:
            html = response.read().decode('utf-8', errors='ignore')
            
        parser = SimpleHTMLScraper()
        parser.feed(html)
        
        name = clean_wine_title(parser.og_title or parser.title)
        description = parser.og_desc or (parser.paragraphs[0] if parser.paragraphs else "")
        photo_url = parser.og_image
        
        lower_html = html.lower()
        
        category = "VINO_ROSSO"
        if "spumante" in lower_html or "brut" in lower_html:
            category = "SPUMANTE"
        elif "rosato" in lower_html or "rosé" in lower_html:
            category = "ROSATO"
        elif "bianco" in lower_html:
            category = "VINO_BIANCO"
        elif "passito" in lower_html:
            category = "PASSITO"
            
        denominazione = "DOC"
        if "tintilia" in lower_html or "biferno" in lower_html or "doc" in lower_html:
            denominazione = "DOC"
        elif "igt" in lower_html:
            denominazione = "IGT"
            
        is_riserva = "riserva" in lower_html
        
        # Catalogo Vini Molise è Senza Annata (S.A.) by default
        vintage_year = None
        
        alc_match = re.search(r'(\d{2}(?:[.,]\d)?)\s*%\s*(?:vol)?', lower_html)
        alcohol_degrees = float(alc_match.group(1).replace(',', '.')) if alc_match else None
        
        return {
            "name": name.strip(),
            "category": category,
            "denominazione": denominazione,
            "vintage_year": vintage_year,
            "is_riserva": is_riserva,
            "alcohol_degrees": alcohol_degrees,
            "description": description.strip(),
            "photo_url": photo_url,
            "technical_sheet_pdf": ""
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore durante l'analisi dell'URL: {str(e)}")

@router.get("", response_model=List[ProductResponse])
async def get_products(
    category: Optional[str] = None,
    denominazione: Optional[str] = None,
    producer_id: Optional[str] = None,
    vintage_year: Optional[int] = None,
    is_riserva: Optional[bool] = None,
    search: Optional[str] = None,
    status: Optional[str] = "PUBLISHED",
    db=Depends(get_database)
):
    query = {}
    if status and status != "ALL":
        query["status"] = status
    if category:
        query["category"] = category
    if denominazione:
        query["denominazione"] = denominazione
    if vintage_year:
        query["vintage_year"] = vintage_year
    if is_riserva is not None:
        query["is_riserva"] = is_riserva
    if producer_id and ObjectId.is_valid(producer_id):
        query["producer_id"] = ObjectId(producer_id)
    if search:
        query["$or"] = [
            {"name": {"$regex": search, "$options": "i"}},
            {"description": {"$regex": search, "$options": "i"}},
            {"grape_varieties": {"$elemMatch": {"$regex": search, "$options": "i"}}}
        ]
        
    cursor = db.products.find(query).sort("created_at", -1)
    raw_products = await cursor.to_list(1000)

    # Bulk fetch all producers to eliminate N+1 queries
    producers = await db.producers.find().to_list(1000)
    producer_map = {str(p["_id"]): p for p in producers}

    products = []
    for doc in raw_products:
        doc["id"] = str(doc["_id"])
        doc["producer_id"] = str(doc["producer_id"])
        prod_obj = producer_map.get(doc["producer_id"])
        if prod_obj:
            doc["producer_name"] = prod_obj.get("company_name", "")
            doc["producer_slug"] = prod_obj.get("slug", "")
        else:
            doc["producer_name"] = ""
            doc["producer_slug"] = ""
        products.append(doc)
    return products

def parse_custom_attributes_from_str(val_str: str) -> list:
    if not val_str or not isinstance(val_str, str):
        return []
    val_str = val_str.strip()
    if not val_str:
        return []
    if val_str.startswith("[") and val_str.endswith("]"):
        try:
            parsed = json.loads(val_str)
            if isinstance(parsed, list):
                return [{"name": str(item.get("name", "")).strip(), "value": str(item.get("value", "")).strip()} 
                        for item in parsed if isinstance(item, dict) and item.get("name") and item.get("value")]
        except Exception:
            pass
    result = []
    pairs = re.split(r'[;\n]+', val_str)
    for p in pairs:
        if ":" in p:
            parts = p.split(":", 1)
            k = parts[0].strip()
            v = parts[1].strip()
            if k and v:
                result.append({"name": k, "value": v})
    return result

@router.get("/export/json")
async def export_products_json(
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")
    
    producers = await db.producers.find().to_list(1000)
    producer_map = {str(p["_id"]): p for p in producers}

    raw_products = await db.products.find().sort("created_at", -1).to_list(1000)
    products = []
    for doc in raw_products:
        doc["id"] = str(doc["_id"])
        doc["producer_id"] = str(doc["producer_id"])
        prod_obj = producer_map.get(doc["producer_id"])
        if prod_obj:
            doc["producer_name"] = prod_obj.get("company_name", "")
            doc["producer_slug"] = prod_obj.get("slug", "")
        else:
            doc["producer_name"] = ""
            doc["producer_slug"] = ""
        products.append(doc)
        
    json_data = json.dumps(products, indent=2, default=str, ensure_ascii=False)
    return Response(
        content=json_data,
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=catalogo_vini.json"}
    )

@router.get("/export/excel")
async def export_products_excel(
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    producers = await db.producers.find().to_list(1000)
    producer_map = {str(p["_id"]): p for p in producers}

    raw_products = await db.products.find().sort("created_at", -1).to_list(1000)
    products = []
    for doc in raw_products:
        doc["id"] = str(doc["_id"])
        doc["producer_id"] = str(doc["producer_id"])
        prod_obj = producer_map.get(doc["producer_id"])
        if prod_obj:
            doc["producer_name"] = prod_obj.get("company_name", "")
            doc["producer_slug"] = prod_obj.get("slug", "")
        else:
            doc["producer_name"] = ""
            doc["producer_slug"] = ""
        products.append(doc)

    # Collect all registered attribute names from db.attributes first
    attr_cursor = db.attributes.find().sort("name", 1)
    attr_names = []
    async for attr_doc in attr_cursor:
        aname = attr_doc.get("name", "").strip()
        if aname and aname not in attr_names:
            attr_names.append(aname)

    # Collect any custom attribute names from products that might not be in db.attributes
    for p in products:
        custom_attrs = p.get("custom_attributes") or []
        for ca in custom_attrs:
            if isinstance(ca, dict) and ca.get("name"):
                cname = ca.get("name").strip()
                if cname and cname not in attr_names:
                    attr_names.append(cname)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalogo Vini"

    base_headers = [
        "ID Prodotto", "Cantina", "ID Cantina", "Nome Vino", "Tipologia",
        "Denominazione", "Annata", "Riserva", "Alcol (% Vol)", "Vitigni",
        "Abbinamenti Gastronomici", "Temperatura Servizio", "Prezzo Indicativo",
        "Descrizione", "Esame Visivo", "Esame Olfattivo", "Esame Gustativo",
        "Foto (URL)", "PDF Scheda (URL)"
    ]

    headers = base_headers + attr_names + ["Stato"]

    ws.append(headers)

    header_fill = PatternFill(start_color="581C26", end_color="581C26", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    ws.row_dimensions[1].height = 28

    for p in products:
        tasting = p.get("tasting_notes") or {}
        custom_attrs = p.get("custom_attributes") or []
        attr_val_map = {}
        for ca in custom_attrs:
            if isinstance(ca, dict) and ca.get("name"):
                attr_val_map[ca["name"].strip()] = ca.get("value", "")

        row_data = [
            p.get("id", ""),
            p.get("producer_name", ""),
            p.get("producer_id", ""),
            p.get("name", ""),
            p.get("category", ""),
            p.get("denominazione", ""),
            p.get("vintage_year") or "",
            "Sì" if p.get("is_riserva") else "No",
            p.get("alcohol_degrees") if p.get("alcohol_degrees") is not None else "",
            ", ".join(p.get("grape_varieties") or []),
            ", ".join(p.get("food_pairings") or []),
            p.get("serving_temperature", ""),
            p.get("indicative_price", ""),
            p.get("description", ""),
            tasting.get("visual", "") if isinstance(tasting, dict) else "",
            tasting.get("olfactory", "") if isinstance(tasting, dict) else "",
            tasting.get("taste", "") if isinstance(tasting, dict) else "",
            ", ".join(p.get("photos") or []),
            p.get("technical_sheet_pdf", "")
        ]

        # Append dynamic attribute column values
        for aname in attr_names:
            row_data.append(attr_val_map.get(aname, ""))

        row_data.append(p.get("status", "PUBLISHED"))
        ws.append(row_data)

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = val_str.split('\n')[0]
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=catalogo_vini.xlsx"}
    )

@router.post("/import/json")
async def import_products_json(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    contents = await file.read()
    try:
        data = json.loads(contents.decode("utf-8"))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"File JSON non valido: {str(e)}")

    if not isinstance(data, list):
        if isinstance(data, dict) and "products" in data:
            data = data["products"]
        else:
            data = [data]

    created = 0
    updated = 0
    errors = 0
    error_details = []

    producers = await db.producers.find().to_list(1000)
    producer_map_name = {p["company_name"].strip().lower(): p["_id"] for p in producers if p.get("company_name")}
    producer_map_id = {str(p["_id"]): p["_id"] for p in producers}
    fallback_producer_id = producers[0]["_id"] if producers else None

    for idx, item in enumerate(data, 1):
        try:
            if not isinstance(item, dict):
                continue

            prod_name = clean_wine_title(item.get("name", "").strip())
            if not prod_name:
                errors += 1
                error_details.append(f"Riga {idx}: Nome vino mancante")
                continue

            p_id = item.get("id") or item.get("_id")
            existing = None
            if p_id and ObjectId.is_valid(str(p_id)):
                existing = await db.products.find_one({"_id": ObjectId(str(p_id))})

            target_producer_id = None
            raw_p_id = item.get("producer_id")
            raw_p_name = item.get("producer_name") or item.get("cantina")
            
            if raw_p_id and str(raw_p_id) in producer_map_id:
                target_producer_id = producer_map_id[str(raw_p_id)]
            elif raw_p_name and str(raw_p_name).strip().lower() in producer_map_name:
                target_producer_id = producer_map_name[str(raw_p_name).strip().lower()]
            elif existing:
                target_producer_id = existing.get("producer_id")
            else:
                target_producer_id = fallback_producer_id

            if not target_producer_id:
                errors += 1
                error_details.append(f"Riga {idx} ({prod_name}): Cantina non trovata")
                continue

            if not existing:
                v_year = item.get("vintage_year")
                q = {"producer_id": target_producer_id, "name": {"$regex": f"^{re.escape(prod_name)}$", "$options": "i"}}
                if v_year:
                    q["vintage_year"] = int(v_year)
                existing = await db.products.find_one(q)

            tasting = item.get("tasting_notes") if isinstance(item.get("tasting_notes"), dict) else {}
            custom_attrs = item.get("custom_attributes") if isinstance(item.get("custom_attributes"), list) else []

            doc = {
                "name": prod_name,
                "producer_id": target_producer_id,
                "category": item.get("category", "VINO_ROSSO"),
                "denominazione": parse_denominazione_acronym(item.get("denominazione", "DOC")),
                "vintage_year": int(item["vintage_year"]) if item.get("vintage_year") and str(item["vintage_year"]).isdigit() else None,
                "is_riserva": bool(item.get("is_riserva", False)),
                "alcohol_degrees": float(item["alcohol_degrees"]) if item.get("alcohol_degrees") is not None and str(item["alcohol_degrees"]).replace('.','',1).isdigit() else None,
                "grape_varieties": item.get("grape_varieties", []) if isinstance(item.get("grape_varieties"), list) else [],
                "food_pairings": item.get("food_pairings", []) if isinstance(item.get("food_pairings"), list) else [],
                "description": item.get("description", ""),
                "tasting_notes": {
                    "visual": tasting.get("visual", ""),
                    "olfactory": tasting.get("olfactory", ""),
                    "taste": tasting.get("taste", "")
                },
                "serving_temperature": item.get("serving_temperature", "16-18°C"),
                "indicative_price": item.get("indicative_price", ""),
                "photos": item.get("photos", []) if isinstance(item.get("photos"), list) else [],
                "technical_sheet_pdf": item.get("technical_sheet_pdf", ""),
                "custom_attributes": custom_attrs,
                "status": item.get("status", "PUBLISHED"),
                "updated_at": datetime.utcnow()
            }

            raw_slug = item.get("slug")
            if raw_slug:
                doc["slug"] = clean_product_slug(raw_slug)
            else:
                doc["slug"] = await generate_unique_product_slug(
                    db=db,
                    name=doc["name"],
                    producer_id=target_producer_id,
                    is_riserva=doc["is_riserva"],
                    exclude_id=existing["_id"] if existing else None
                )

            if existing:
                await db.products.update_one({"_id": existing["_id"]}, {"$set": doc})
                updated += 1
            else:
                doc["created_at"] = datetime.utcnow()
                res = await db.products.insert_one(doc)
                created += 1

            await sync_custom_attributes_with_master(custom_attrs, db)
            await sync_grapes_with_master(doc["grape_varieties"], db)
            await sync_pairings_with_master(doc["food_pairings"], db)

        except Exception as e:
            errors += 1
            error_details.append(f"Riga {idx}: {str(e)}")

    return {
        "message": "Importazione JSON completata",
        "created": created,
        "updated": updated,
        "errors": errors,
        "error_details": error_details[:10]
    }

@router.post("/import/excel")
async def import_products_excel(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"File Excel non valido: {str(e)}")

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="Il file Excel non contiene righe di dati")

    raw_headers_orig = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    raw_headers = [h.lower() for h in raw_headers_orig]
    
    def get_col_idx(names: List[str]) -> Optional[int]:
        for n in names:
            for idx, h in enumerate(raw_headers):
                if n in h:
                    return idx
        return None

    id_idx = get_col_idx(["id prodotto", "_id", "id_prodotto"])
    producer_name_idx = get_col_idx(["cantina", "producer_name", "nome cantina"])
    producer_id_idx = get_col_idx(["id cantina", "producer_id"])
    name_idx = get_col_idx(["nome vino", "name", "nome"])
    cat_idx = get_col_idx(["tipologia", "category"])
    denom_idx = get_col_idx(["denominazione"])
    vintage_idx = get_col_idx(["annata", "vintage_year"])
    riserva_idx = get_col_idx(["riserva", "is_riserva"])
    alcohol_idx = get_col_idx(["alcol", "alcohol_degrees"])
    grapes_idx = get_col_idx(["vitigni", "grape_varieties"])
    pairings_idx = get_col_idx(["abbinamenti", "food_pairings"])
    temp_idx = get_col_idx(["temperatura", "serving_temperature"])
    price_idx = get_col_idx(["prezzo", "indicative_price"])
    desc_idx = get_col_idx(["descrizione", "description"])
    visual_idx = get_col_idx(["visivo", "tasting_visual"])
    olfactory_idx = get_col_idx(["olfattivo", "tasting_olfactory"])
    taste_idx = get_col_idx(["gustativo", "tasting_taste"])
    photos_idx = get_col_idx(["foto", "photos", "url foto"])
    pdf_idx = get_col_idx(["pdf", "technical_sheet_pdf"])
    custom_idx = get_col_idx(["attributi personalizzati", "custom_attributes"])
    status_idx = get_col_idx(["stato", "status"])

    standard_indices = {
        id_idx, producer_name_idx, producer_id_idx, name_idx, cat_idx, denom_idx,
        vintage_idx, riserva_idx, alcohol_idx, grapes_idx, pairings_idx, temp_idx,
        price_idx, desc_idx, visual_idx, olfactory_idx, taste_idx, photos_idx,
        pdf_idx, custom_idx, status_idx
    }
    standard_indices.discard(None)

    producers = await db.producers.find().to_list(1000)
    producer_map_name = {p["company_name"].strip().lower(): p["_id"] for p in producers if p.get("company_name")}
    producer_map_id = {str(p["_id"]): p["_id"] for p in producers}
    fallback_producer_id = producers[0]["_id"] if producers else None

    existing_attrs_docs = await db.attributes.find().to_list(1000)
    attr_cache = {a["name"].strip().lower(): a for a in existing_attrs_docs if a.get("name")}

    existing_grapes_docs = await db.grapes.find().to_list(1000)
    grapes_cache = {g["name"].strip().lower(): g for g in existing_grapes_docs if g.get("name")}

    existing_pairings_docs = await db.pairings.find().to_list(1000)
    pairings_cache = {p["name"].strip().lower(): p for p in existing_pairings_docs if p.get("name")}

    created = 0
    updated = 0
    errors = 0
    error_details = []

    for r_idx, row in enumerate(rows[1:], 2):
        try:
            raw_p_name = str(row[name_idx]).strip() if name_idx is not None and row[name_idx] is not None else ""
            prod_name = clean_wine_title(raw_p_name)
            if not prod_name:
                continue

            p_id = str(row[id_idx]).strip() if id_idx is not None and row[id_idx] is not None else None
            existing = None
            if p_id and ObjectId.is_valid(p_id):
                existing = await db.products.find_one({"_id": ObjectId(p_id)})

            raw_p_id = str(row[producer_id_idx]).strip() if producer_id_idx is not None and row[producer_id_idx] is not None else None
            raw_p_name = str(row[producer_name_idx]).strip() if producer_name_idx is not None and row[producer_name_idx] is not None else None

            target_producer_id = None
            if raw_p_id and raw_p_id in producer_map_id:
                target_producer_id = producer_map_id[raw_p_id]
            elif raw_p_name and raw_p_name.lower() in producer_map_name:
                target_producer_id = producer_map_name[raw_p_name.lower()]
            elif existing:
                target_producer_id = existing.get("producer_id")
            else:
                target_producer_id = fallback_producer_id

            if not target_producer_id:
                errors += 1
                error_details.append(f"Riga {r_idx} ({prod_name}): Cantina non trovata")
                continue

            v_year = None
            if vintage_idx is not None and row[vintage_idx] is not None:
                v_str = re.sub(r'\D', '', str(row[vintage_idx]))
                if v_str and len(v_str) == 4:
                    v_year = int(v_str)

            if not existing:
                q = {"producer_id": target_producer_id, "name": {"$regex": f"^{re.escape(prod_name)}$", "$options": "i"}}
                if v_year:
                    q["vintage_year"] = v_year
                existing = await db.products.find_one(q)

            is_ris = False
            if riserva_idx is not None and row[riserva_idx] is not None:
                r_val = str(row[riserva_idx]).strip().lower()
                is_ris = r_val in ["sì", "si", "true", "1", "riserva"]

            alc = None
            if alcohol_idx is not None and row[alcohol_idx] is not None:
                alc_str = str(row[alcohol_idx]).replace(',', '.')
                m_alc = re.search(r'(\d+(?:\.\d+)?)', alc_str)
                if m_alc:
                    alc = float(m_alc.group(1))

            grapes = [g.strip() for g in str(row[grapes_idx]).split(',') if g.strip()] if grapes_idx is not None and row[grapes_idx] is not None else []
            pairings = [p.strip() for p in str(row[pairings_idx]).split(',') if p.strip()] if pairings_idx is not None and row[pairings_idx] is not None else []
            photos = [ph.strip() for ph in str(row[photos_idx]).split(',') if ph.strip()] if photos_idx is not None and row[photos_idx] is not None else []

            custom_str = str(row[custom_idx]) if custom_idx is not None and row[custom_idx] is not None else ""
            custom_attrs = parse_custom_attributes_from_str(custom_str)

            # Process individual dynamic attribute columns
            for col_i, header_orig in enumerate(raw_headers_orig):
                if col_i in standard_indices:
                    continue
                if not header_orig:
                    continue
                if col_i < len(row) and row[col_i] is not None:
                    cell_val = str(row[col_i]).strip()
                    if cell_val:
                        existing_attr = next((a for a in custom_attrs if isinstance(a, dict) and a.get("name", "").lower() == header_orig.lower()), None)
                        if existing_attr:
                            existing_attr["value"] = cell_val
                        else:
                            custom_attrs.append({"name": header_orig, "value": cell_val})

            doc = {
                "name": prod_name,
                "producer_id": target_producer_id,
                "category": str(row[cat_idx]).strip().upper() if cat_idx is not None and row[cat_idx] else "VINO_ROSSO",
                "denominazione": parse_denominazione_acronym(str(row[denom_idx])) if denom_idx is not None and row[denom_idx] else "DOC",
                "vintage_year": v_year,
                "is_riserva": is_ris,
                "alcohol_degrees": alc,
                "grape_varieties": grapes,
                "food_pairings": pairings,
                "description": str(row[desc_idx]).strip() if desc_idx is not None and row[desc_idx] is not None else "",
                "tasting_notes": {
                    "visual": str(row[visual_idx]).strip() if visual_idx is not None and row[visual_idx] is not None else "",
                    "olfactory": str(row[olfactory_idx]).strip() if olfactory_idx is not None and row[olfactory_idx] is not None else "",
                    "taste": str(row[taste_idx]).strip() if taste_idx is not None and row[taste_idx] is not None else ""
                },
                "serving_temperature": str(row[temp_idx]).strip() if temp_idx is not None and row[temp_idx] is not None else "16-18°C",
                "indicative_price": str(row[price_idx]).strip() if price_idx is not None and row[price_idx] is not None else "",
                "photos": photos,
                "technical_sheet_pdf": str(row[pdf_idx]).strip() if pdf_idx is not None and row[pdf_idx] is not None else "",
                "custom_attributes": custom_attrs,
                "status": str(row[status_idx]).strip().upper() if status_idx is not None and row[status_idx] else "PUBLISHED",
                "updated_at": datetime.utcnow()
            }

            doc["slug"] = await generate_unique_product_slug(
                db=db,
                name=doc["name"],
                producer_id=target_producer_id,
                is_riserva=doc["is_riserva"],
                exclude_id=existing["_id"] if existing else None
            )

            if existing:
                await db.products.update_one({"_id": existing["_id"]}, {"$set": doc})
                updated += 1
            else:
                doc["created_at"] = datetime.utcnow()
                res = await db.products.insert_one(doc)
                created += 1

            await sync_custom_attributes_with_master(custom_attrs, db, attr_cache)
            await sync_grapes_with_master(grapes, db, grapes_cache)
            await sync_pairings_with_master(pairings, db, pairings_cache)

        except Exception as e:
            errors += 1
            error_details.append(f"Riga {r_idx}: {str(e)}")

    return {
        "message": "Importazione Excel completata",
        "created": created,
        "updated": updated,
        "errors": errors,
        "error_details": error_details[:10]
    }

@router.get("/{identifier}", response_model=ProductResponse)
async def get_product_by_slug_or_id(identifier: str, db=Depends(get_database)):
    query = {"$or": [{"slug": identifier}]}
    cleaned = clean_product_slug(identifier)
    if cleaned and cleaned != identifier:
        query["$or"].append({"slug": cleaned})
    if ObjectId.is_valid(identifier):
        query["$or"].append({"_id": ObjectId(identifier)})
        
    doc = await db.products.find_one(query)
    if not doc:
        # Fallback flessibile: ad es. se viene cercato "tintilia-molise", trova "tintilia-del-molise"
        parts = [re.escape(p) for p in identifier.split('-') if p and p not in ['del', 'di', 'dei', 'della', 'degli', 'doc', 'igt', 'dop', 'docg', 'igp']]
        if parts:
            regex_pattern = ".*".join(parts)
            doc = await db.products.find_one({"slug": {"$regex": f"^{regex_pattern}", "$options": "i"}})

    if not doc:
        raise HTTPException(status_code=404, detail="Prodotto non trovato")
        
    return await format_product_response(doc, db)

async def sync_custom_attributes_with_master(custom_attributes: list, db, attr_cache: dict = None):
    if not custom_attributes:
        return
    for attr in custom_attributes:
        name = attr.get("name", "").strip() if isinstance(attr, dict) else (getattr(attr, "name", "") or "").strip()
        val = attr.get("value", "").strip() if isinstance(attr, dict) else (getattr(attr, "value", "") or "").strip()
        if not name or not val:
            continue
        name_lower = name.lower()
        existing = None
        if attr_cache is not None:
            existing = attr_cache.get(name_lower)
        else:
            existing = await db.attributes.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}})
        
        if existing:
            current_suggs = existing.get("suggested_values", [])
            if not any(s.strip().lower() == val.strip().lower() for s in current_suggs):
                current_suggs.append(val.strip())
                existing["suggested_values"] = current_suggs
                await db.attributes.update_one(
                    {"_id": existing["_id"]},
                    {"$addToSet": {"suggested_values": val.strip()}}
                )
        else:
            new_doc = {
                "name": name,
                "unit_or_hint": "",
                "suggested_values": [val.strip()],
                "created_at": datetime.utcnow()
            }
            res = await db.attributes.insert_one(new_doc)
            new_doc["_id"] = res.inserted_id
            if attr_cache is not None:
                attr_cache[name_lower] = new_doc

async def sync_grapes_with_master(grape_varieties: list, db, grapes_cache: dict = None):
    if not grape_varieties:
        return
    for item in grape_varieties:
        if not item or not isinstance(item, str):
            continue
        clean_name = re.sub(r'\d+\s*%?', '', item).strip()
        if not clean_name:
            continue
        clean_lower = clean_name.lower()
        existing = None
        if grapes_cache is not None:
            existing = grapes_cache.get(clean_lower)
        else:
            existing = await db.grapes.find_one({"name": {"$regex": f"^{re.escape(clean_name)}$", "$options": "i"}})
        if not existing:
            new_doc = {
                "name": clean_name,
                "category": "AUTOCTONO",
                "created_at": datetime.utcnow()
            }
            res = await db.grapes.insert_one(new_doc)
            new_doc["_id"] = res.inserted_id
            if grapes_cache is not None:
                grapes_cache[clean_lower] = new_doc

async def sync_pairings_with_master(food_pairings: list, db, pairings_cache: dict = None):
    if not food_pairings:
        return
    for item in food_pairings:
        if not item or not isinstance(item, str):
            continue
        clean_name = item.strip()
        if not clean_name:
            continue
        clean_lower = clean_name.lower()
        existing = None
        if pairings_cache is not None:
            existing = pairings_cache.get(clean_lower)
        else:
            existing = await db.pairings.find_one({"name": {"$regex": f"^{re.escape(clean_name)}$", "$options": "i"}})
        if not existing:
            new_doc = {
                "name": clean_name,
                "category": "GENERALE",
                "created_at": datetime.utcnow()
            }
            res = await db.pairings.insert_one(new_doc)
            new_doc["_id"] = res.inserted_id
            if pairings_cache is not None:
                pairings_cache[clean_lower] = new_doc

@router.post("", response_model=ProductResponse)
async def create_product(
    product_in: ProductCreate,
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    is_admin = current_user.get("role") == "ADMIN"
    
    if is_admin:
        if not product_in.producer_id or not ObjectId.is_valid(product_in.producer_id):
            raise HTTPException(status_code=400, detail="Devi specificare una cantina valida per questo prodotto")
        target_producer_id = ObjectId(product_in.producer_id)
    else:
        # Producer user
        user_producer_id = current_user.get("producer_id")
        if not user_producer_id or not ObjectId.is_valid(user_producer_id):
            raise HTTPException(status_code=400, detail="Il tuo account non è collegato ad alcuna cantina")
        target_producer_id = ObjectId(user_producer_id)
        
    # Verify producer exists
    producer = await db.producers.find_one({"_id": target_producer_id})
    if not producer:
        raise HTTPException(status_code=404, detail="Cantina non trovata")
        
    cleaned_name = clean_wine_title(product_in.name)
    if product_in.slug:
        slug = clean_product_slug(product_in.slug)
    else:
        slug = await generate_unique_product_slug(
            db=db,
            name=cleaned_name,
            producer_id=target_producer_id,
            is_riserva=product_in.is_riserva
        )
        
    doc = product_in.model_dump()
    doc["name"] = cleaned_name
    doc["producer_id"] = target_producer_id
    doc["denominazione"] = parse_denominazione_acronym(doc.get("denominazione", "DOC"))
    doc["slug"] = slug
    doc["created_at"] = datetime.utcnow()
    doc["updated_at"] = datetime.utcnow()
    
    res = await db.products.insert_one(doc)
    doc["_id"] = res.inserted_id

    # Auto-sync custom attributes, grapes & food pairings with master collections
    await sync_custom_attributes_with_master(doc.get("custom_attributes", []), db)
    await sync_grapes_with_master(doc.get("grape_varieties", []), db)
    await sync_pairings_with_master(doc.get("food_pairings", []), db)

    return await format_product_response(doc, db)

@router.post("/{product_id}/clone", response_model=ProductResponse)
async def clone_product(
    product_id: str,
    target_producer_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if not ObjectId.is_valid(product_id):
        raise HTTPException(status_code=400, detail="ID prodotto non valido")
        
    original = await db.products.find_one({"_id": ObjectId(product_id)})
    if not original:
        raise HTTPException(status_code=404, detail="Prodotto da clonare non trovato")
        
    is_admin = current_user.get("role") == "ADMIN"
    user_producer_id = current_user.get("producer_id")
    
    # Check permissions
    if not is_admin and str(original["producer_id"]) != str(user_producer_id):
        raise HTTPException(status_code=403, detail="Non puoi clonare i prodotti di un'altra cantina")
        
    # Determine new producer_id
    if is_admin and target_producer_id and ObjectId.is_valid(target_producer_id):
        new_producer_id = ObjectId(target_producer_id)
    else:
        new_producer_id = original["producer_id"]
        
    cloned_doc = dict(original)
    del cloned_doc["_id"]
    
    cloned_name = f"{clean_wine_title(original['name'])} (Copia)"
    cloned_doc["name"] = cloned_name
    cloned_doc["slug"] = f"{clean_product_slug(cloned_name)}-{int(datetime.utcnow().timestamp())}"
    cloned_doc["vintage_year"] = None # Reset to Senza Annata (S.A.)
    cloned_doc["status"] = "DRAFT" # Draft until user edits
    cloned_doc["created_at"] = datetime.utcnow()
    cloned_doc["updated_at"] = datetime.utcnow()
    
    res = await db.products.insert_one(cloned_doc)
    cloned_doc["_id"] = res.inserted_id
    return await format_product_response(cloned_doc, db)

@router.put("/{product_id}", response_model=ProductResponse)
async def update_product(
    product_id: str,
    product_in: ProductUpdate,
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if not ObjectId.is_valid(product_id):
        raise HTTPException(status_code=400, detail="ID prodotto non valido")
        
    existing = await db.products.find_one({"_id": ObjectId(product_id)})
    if not existing:
        raise HTTPException(status_code=404, detail="Prodotto non trovato")
        
    is_admin = current_user.get("role") == "ADMIN"
    user_producer_id = current_user.get("producer_id")
    
    if not is_admin and str(existing["producer_id"]) != str(user_producer_id):
        raise HTTPException(status_code=403, detail="Non puoi modificare questo prodotto")
        
    update_data = {k: v for k, v in product_in.model_dump().items() if v is not None}
    
    # Handle producer_id reassignment by Admin
    if "producer_id" in update_data:
        if is_admin and ObjectId.is_valid(update_data["producer_id"]):
            update_data["producer_id"] = ObjectId(update_data["producer_id"])
        else:
            del update_data["producer_id"]
            
    if "denominazione" in update_data:
        update_data["denominazione"] = parse_denominazione_acronym(update_data["denominazione"])

    if "name" in update_data:
        update_data["name"] = clean_wine_title(update_data["name"])
        if update_data["name"] != existing.get("name"):
            is_ris = update_data.get('is_riserva', existing.get('is_riserva', False))
            target_prod_id = update_data.get("producer_id", existing["producer_id"])
            update_data["slug"] = await generate_unique_product_slug(
                db=db,
                name=update_data["name"],
                producer_id=target_prod_id,
                is_riserva=is_ris,
                exclude_id=existing["_id"]
            )

    if "slug" in update_data:
        update_data["slug"] = clean_product_slug(update_data["slug"])
        
    update_data["updated_at"] = datetime.utcnow()
    
    await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": update_data})
    
    if "custom_attributes" in update_data:
        await sync_custom_attributes_with_master(update_data["custom_attributes"], db)
    if "grape_varieties" in update_data:
        await sync_grapes_with_master(update_data["grape_varieties"], db)
    if "food_pairings" in update_data:
        await sync_pairings_with_master(update_data["food_pairings"], db)

    updated = await db.products.find_one({"_id": ObjectId(product_id)})
    return await format_product_response(updated, db)

@router.delete("/{product_id}")
async def delete_product(
    product_id: str,
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if not ObjectId.is_valid(product_id):
        raise HTTPException(status_code=400, detail="ID prodotto non valido")
        
    existing = await db.products.find_one({"_id": ObjectId(product_id)})
    if not existing:
        raise HTTPException(status_code=404, detail="Prodotto non trovato")
        
    is_admin = current_user.get("role") == "ADMIN"
    user_producer_id = current_user.get("producer_id")
    
    if not is_admin and str(existing["producer_id"]) != str(user_producer_id):
        raise HTTPException(status_code=403, detail="Non puoi eliminare questo prodotto")
        
    await db.products.delete_one({"_id": ObjectId(product_id)})
    return {"message": "Prodotto eliminato con successo"}

async def run_products_cleanup_migration(db) -> int:
    products = await db.products.find({}).to_list(length=10000)
    migrated_count = 0
    for prod in products:
        p_id = prod["_id"]
        old_name = prod.get("name", "")
        old_slug = prod.get("slug", "")
        old_vintage = prod.get("vintage_year")
        old_denom = prod.get("denominazione", "")
        
        # 1. Clean wine title from dates and denominations
        new_name = clean_wine_title(old_name)
        
        # 2. Standardize denominazione acronym
        new_denom = parse_denominazione_acronym(old_denom) if old_denom else "DOC"
        
        # 3. Always set vintage_year = None (Catalogo Senza Annata - S.A.)
        new_vintage = None
        
        # 4. Generate unique clean slug without vintage_year and without denominations
        new_slug = await generate_unique_product_slug(
            db=db,
            name=new_name,
            producer_id=prod.get("producer_id"),
            is_riserva=bool(prod.get("is_riserva", False)),
            exclude_id=p_id
        )
        
        update_fields = {}
        if new_name != old_name:
            update_fields["name"] = new_name
        if new_denom != old_denom:
            update_fields["denominazione"] = new_denom
        if old_vintage is not None:
            update_fields["vintage_year"] = new_vintage
        if new_slug != old_slug:
            update_fields["slug"] = new_slug
            
        if update_fields:
            update_fields["updated_at"] = datetime.utcnow()
            await db.products.update_one({"_id": p_id}, {"$set": update_fields})
            migrated_count += 1
            
    return migrated_count

@router.post("/cleanup-slugs-and-titles")
async def cleanup_slugs_and_titles(
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Solo gli amministratori possono eseguire la pulizia")
    
    migrated_count = await run_products_cleanup_migration(db)
    return {"message": f"Pulizia catalogo completata con successo: {migrated_count} vini aggiornati a Senza Annata (S.A.) e slug puliti."}
