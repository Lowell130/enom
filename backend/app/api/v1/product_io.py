"""Import/export del catalogo (JSON, Excel). L'import da PDF e' in pdf_import.py.

Questo router e' montato su /api/v1/products PRIMA del router principale dei prodotti,
cosi' percorsi come /export/json non vengono intercettati da GET /{identifier}."""
import io
import json
import re
from datetime import datetime
from typing import List, Optional

import openpyxl
from bson import ObjectId
from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.api.v1.auth import get_current_user
from app.db.mongodb import get_database
from app.services.catalog import (
    clean_wine_title,
    ensure_unique_product_slug,
    generate_unique_product_slug,
    parse_denominazione_acronym,
)
from app.services.taxonomy import (
    sync_custom_attributes_with_master,
    sync_grapes_with_master,
    sync_pairings_with_master,
)

router = APIRouter()


def parse_custom_attributes_from_str(val_str: str) -> list:
    if not val_str or not isinstance(val_str, str):
        return []
    val_str = val_str.strip()
    if not val_str:
        return []
    if val_str.startswith("[") and val_str.endswith("]"):
        try:
            parsed = json.loads(val_str)
            if isinstance(parsed, list):
                return [{"name": str(item.get("name", "")).strip(), "value": str(item.get("value", "")).strip()} 
                        for item in parsed if isinstance(item, dict) and item.get("name") and item.get("value")]
        except Exception:
            pass
    result = []
    pairs = re.split(r'[;\n]+', val_str)
    for p in pairs:
        if ":" in p:
            parts = p.split(":", 1)
            k = parts[0].strip()
            v = parts[1].strip()
            if k and v:
                result.append({"name": k, "value": v})
    return result

ILLEGAL_CHARACTERS_RE = re.compile(r'[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F-\x84\x86-\x9F]')

def clean_excel_val(val):
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    val_str = str(val)
    return ILLEGAL_CHARACTERS_RE.sub("", val_str)

@router.get("/export/json")
async def export_products_json(
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")
    
    producers = await db.producers.find().to_list(1000)
    producer_map = {str(p["_id"]): p for p in producers}

    raw_products = await db.products.find().sort("created_at", -1).to_list(1000)
    products = []
    for doc in raw_products:
        doc["id"] = str(doc["_id"])
        doc["producer_id"] = str(doc["producer_id"])
        prod_obj = producer_map.get(doc["producer_id"])
        if prod_obj:
            doc["producer_name"] = prod_obj.get("company_name", "")
            doc["producer_slug"] = prod_obj.get("slug", "")
        else:
            doc["producer_name"] = ""
            doc["producer_slug"] = ""
        products.append(doc)
        
    json_data = json.dumps(products, indent=2, default=str, ensure_ascii=False)
    return Response(
        content=json_data,
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=catalogo_vini.json"}
    )

@router.get("/export/excel")
async def export_products_excel(
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    producers = await db.producers.find().to_list(1000)
    producer_map = {str(p["_id"]): p for p in producers}

    raw_products = await db.products.find().sort("created_at", -1).to_list(1000)
    products = []
    for doc in raw_products:
        doc["id"] = str(doc["_id"])
        doc["producer_id"] = str(doc.get("producer_id", ""))
        prod_obj = producer_map.get(doc["producer_id"])
        if prod_obj:
            doc["producer_name"] = prod_obj.get("company_name", "")
            doc["producer_slug"] = prod_obj.get("slug", "")
        else:
            doc["producer_name"] = ""
            doc["producer_slug"] = ""
        products.append(doc)

    # Collect all registered attribute names from db.attributes first
    attr_cursor = db.attributes.find().sort("name", 1)
    attr_names = []
    async for attr_doc in attr_cursor:
        aname = attr_doc.get("name", "").strip()
        if aname and aname not in attr_names:
            attr_names.append(aname)

    # Collect any custom attribute names from products that might not be in db.attributes
    for p in products:
        custom_attrs = p.get("custom_attributes") or []
        for ca in custom_attrs:
            if isinstance(ca, dict) and ca.get("name"):
                cname = ca.get("name").strip()
                if cname and cname not in attr_names:
                    attr_names.append(cname)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalogo Vini"

    base_headers = [
        "ID Prodotto", "Cantina", "ID Cantina", "Nome Vino", "Tipologia",
        "Denominazione", "Annata", "Riserva", "Alcol (% Vol)", "Vitigni",
        "Abbinamenti Gastronomici", "Temperatura Servizio", "Prezzo Indicativo",
        "Descrizione", "Esame Visivo", "Esame Olfattivo", "Esame Gustativo",
        "Foto (URL)", "PDF Scheda (URL)"
    ]

    headers = base_headers + attr_names + ["Stato"]

    ws.append([clean_excel_val(h) for h in headers])

    header_fill = PatternFill(start_color="581C26", end_color="581C26", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    ws.row_dimensions[1].height = 28

    for p in products:
        tasting = p.get("tasting_notes") or {}
        custom_attrs = p.get("custom_attributes") or []
        attr_val_map = {}
        for ca in custom_attrs:
            if isinstance(ca, dict) and ca.get("name"):
                attr_val_map[ca["name"].strip()] = ca.get("value", "")

        row_data = [
            p.get("id", ""),
            p.get("producer_name", ""),
            p.get("producer_id", ""),
            p.get("name", ""),
            p.get("category", ""),
            p.get("denominazione", ""),
            p.get("vintage_year") or "",
            "Sì" if p.get("is_riserva") else "No",
            p.get("alcohol_degrees") if p.get("alcohol_degrees") is not None else "",
            ", ".join(p.get("grape_varieties") or []),
            ", ".join(p.get("food_pairings") or []),
            p.get("serving_temperature", ""),
            p.get("indicative_price", ""),
            p.get("description", ""),
            tasting.get("visual", "") if isinstance(tasting, dict) else "",
            tasting.get("olfactory", "") if isinstance(tasting, dict) else "",
            tasting.get("taste", "") if isinstance(tasting, dict) else "",
            ", ".join(p.get("photos") or []),
            p.get("technical_sheet_pdf", "")
        ]

        # Append dynamic attribute column values
        for aname in attr_names:
            row_data.append(attr_val_map.get(aname, ""))

        row_data.append(p.get("status", "PUBLISHED"))
        
        clean_row = [clean_excel_val(v) for v in row_data]
        ws.append(clean_row)

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = val_str.split('\n')[0]
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=catalogo_vini.xlsx"}
    )

@router.post("/import/json")
async def import_products_json(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    contents = await file.read()
    try:
        data = json.loads(contents.decode("utf-8"))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"File JSON non valido: {str(e)}")

    if not isinstance(data, list):
        if isinstance(data, dict) and "products" in data:
            data = data["products"]
        else:
            data = [data]

    created = 0
    updated = 0
    errors = 0
    error_details = []

    producers = await db.producers.find().to_list(1000)
    producer_map_name = {p["company_name"].strip().lower(): p["_id"] for p in producers if p.get("company_name")}
    producer_map_id = {str(p["_id"]): p["_id"] for p in producers}
    fallback_producer_id = producers[0]["_id"] if producers else None

    for idx, item in enumerate(data, 1):
        try:
            if not isinstance(item, dict):
                continue

            prod_name = clean_wine_title(item.get("name", "").strip())
            if not prod_name:
                errors += 1
                error_details.append(f"Riga {idx}: Nome vino mancante")
                continue

            p_id = item.get("id") or item.get("_id")
            existing = None
            if p_id and ObjectId.is_valid(str(p_id)):
                existing = await db.products.find_one({"_id": ObjectId(str(p_id))})

            target_producer_id = None
            raw_p_id = item.get("producer_id")
            raw_p_name = item.get("producer_name") or item.get("cantina")
            
            if raw_p_id and str(raw_p_id) in producer_map_id:
                target_producer_id = producer_map_id[str(raw_p_id)]
            elif raw_p_name and str(raw_p_name).strip().lower() in producer_map_name:
                target_producer_id = producer_map_name[str(raw_p_name).strip().lower()]
            elif existing:
                target_producer_id = existing.get("producer_id")
            else:
                target_producer_id = fallback_producer_id

            if not target_producer_id:
                errors += 1
                error_details.append(f"Riga {idx} ({prod_name}): Cantina non trovata")
                continue

            if not existing:
                v_year = item.get("vintage_year")
                q = {"producer_id": target_producer_id, "name": {"$regex": f"^{re.escape(prod_name)}$", "$options": "i"}}
                if v_year:
                    q["vintage_year"] = int(v_year)
                existing = await db.products.find_one(q)

            tasting = item.get("tasting_notes") if isinstance(item.get("tasting_notes"), dict) else {}
            custom_attrs = item.get("custom_attributes") if isinstance(item.get("custom_attributes"), list) else []

            doc = {
                "name": prod_name,
                "producer_id": target_producer_id,
                "category": item.get("category", "VINO_ROSSO"),
                "denominazione": parse_denominazione_acronym(item.get("denominazione", "DOC")),
                "vintage_year": int(item["vintage_year"]) if item.get("vintage_year") and str(item["vintage_year"]).isdigit() else None,
                "is_riserva": bool(item.get("is_riserva", False)),
                "alcohol_degrees": float(item["alcohol_degrees"]) if item.get("alcohol_degrees") is not None and str(item["alcohol_degrees"]).replace('.','',1).isdigit() else None,
                "grape_varieties": item.get("grape_varieties", []) if isinstance(item.get("grape_varieties"), list) else [],
                "food_pairings": item.get("food_pairings", []) if isinstance(item.get("food_pairings"), list) else [],
                "description": item.get("description", ""),
                "tasting_notes": {
                    "visual": tasting.get("visual", ""),
                    "olfactory": tasting.get("olfactory", ""),
                    "taste": tasting.get("taste", "")
                },
                "serving_temperature": item.get("serving_temperature", "16-18°C"),
                "indicative_price": item.get("indicative_price", ""),
                "photos": item.get("photos", []) if isinstance(item.get("photos"), list) else [],
                "technical_sheet_pdf": item.get("technical_sheet_pdf", ""),
                "custom_attributes": custom_attrs,
                "status": "DRAFT" if str(item.get("status", "PUBLISHED")).upper() == "DRAFT" else "PUBLISHED",
                "updated_at": datetime.utcnow()
            }

            raw_slug = item.get("slug")
            if raw_slug:
                doc["slug"] = await ensure_unique_product_slug(
                    db, str(raw_slug), exclude_id=existing["_id"] if existing else None
                )
            else:
                doc["slug"] = await generate_unique_product_slug(
                    db=db,
                    name=doc["name"],
                    producer_id=target_producer_id,
                    is_riserva=doc["is_riserva"],
                    exclude_id=existing["_id"] if existing else None
                )

            if existing:
                await db.products.update_one({"_id": existing["_id"]}, {"$set": doc})
                updated += 1
            else:
                doc["created_at"] = datetime.utcnow()
                res = await db.products.insert_one(doc)
                created += 1

            await sync_custom_attributes_with_master(custom_attrs, db)
            await sync_grapes_with_master(doc["grape_varieties"], db)
            await sync_pairings_with_master(doc["food_pairings"], db)

        except Exception as e:
            errors += 1
            error_details.append(f"Riga {idx}: {str(e)}")

    return {
        "message": "Importazione JSON completata",
        "created": created,
        "updated": updated,
        "errors": errors,
        "error_details": error_details[:10]
    }

@router.post("/import/excel")
async def import_products_excel(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db=Depends(get_database)
):
    if current_user.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="Accesso riservato all'amministratore")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"File Excel non valido: {str(e)}")

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(status_code=400, detail="Il file Excel non contiene righe di dati")

    raw_headers_orig = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    raw_headers = [h.lower() for h in raw_headers_orig]
    
    def get_col_idx(names: List[str]) -> Optional[int]:
        for n in names:
            for idx, h in enumerate(raw_headers):
                if n in h:
                    return idx
        return None

    id_idx = get_col_idx(["id prodotto", "_id", "id_prodotto"])
    producer_name_idx = get_col_idx(["cantina", "producer_name", "nome cantina"])
    producer_id_idx = get_col_idx(["id cantina", "producer_id"])
    name_idx = get_col_idx(["nome vino", "name", "nome"])
    cat_idx = get_col_idx(["tipologia", "category"])
    denom_idx = get_col_idx(["denominazione"])
    vintage_idx = get_col_idx(["annata", "vintage_year"])
    riserva_idx = get_col_idx(["riserva", "is_riserva"])
    alcohol_idx = get_col_idx(["alcol", "alcohol_degrees"])
    grapes_idx = get_col_idx(["vitigni", "grape_varieties"])
    pairings_idx = get_col_idx(["abbinamenti", "food_pairings"])
    temp_idx = get_col_idx(["temperatura", "serving_temperature"])
    price_idx = get_col_idx(["prezzo", "indicative_price"])
    desc_idx = get_col_idx(["descrizione", "description"])
    visual_idx = get_col_idx(["visivo", "tasting_visual"])
    olfactory_idx = get_col_idx(["olfattivo", "tasting_olfactory"])
    taste_idx = get_col_idx(["gustativo", "tasting_taste"])
    photos_idx = get_col_idx(["foto", "photos", "url foto"])
    pdf_idx = get_col_idx(["pdf", "technical_sheet_pdf"])
    custom_idx = get_col_idx(["attributi personalizzati", "custom_attributes"])
    status_idx = get_col_idx(["stato", "status"])

    standard_indices = {
        id_idx, producer_name_idx, producer_id_idx, name_idx, cat_idx, denom_idx,
        vintage_idx, riserva_idx, alcohol_idx, grapes_idx, pairings_idx, temp_idx,
        price_idx, desc_idx, visual_idx, olfactory_idx, taste_idx, photos_idx,
        pdf_idx, custom_idx, status_idx
    }
    standard_indices.discard(None)

    producers = await db.producers.find().to_list(1000)
    producer_map_name = {p["company_name"].strip().lower(): p["_id"] for p in producers if p.get("company_name")}
    producer_map_id = {str(p["_id"]): p["_id"] for p in producers}
    fallback_producer_id = producers[0]["_id"] if producers else None

    existing_attrs_docs = await db.attributes.find().to_list(1000)
    attr_cache = {a["name"].strip().lower(): a for a in existing_attrs_docs if a.get("name")}

    existing_grapes_docs = await db.grapes.find().to_list(1000)
    grapes_cache = {g["name"].strip().lower(): g for g in existing_grapes_docs if g.get("name")}

    existing_pairings_docs = await db.pairings.find().to_list(1000)
    pairings_cache = {p["name"].strip().lower(): p for p in existing_pairings_docs if p.get("name")}

    created = 0
    updated = 0
    errors = 0
    error_details = []

    for r_idx, row in enumerate(rows[1:], 2):
        try:
            raw_p_name = str(row[name_idx]).strip() if name_idx is not None and row[name_idx] is not None else ""
            prod_name = clean_wine_title(raw_p_name)
            if not prod_name:
                continue

            p_id = str(row[id_idx]).strip() if id_idx is not None and row[id_idx] is not None else None
            existing = None
            if p_id and ObjectId.is_valid(p_id):
                existing = await db.products.find_one({"_id": ObjectId(p_id)})

            raw_p_id = str(row[producer_id_idx]).strip() if producer_id_idx is not None and row[producer_id_idx] is not None else None
            raw_p_name = str(row[producer_name_idx]).strip() if producer_name_idx is not None and row[producer_name_idx] is not None else None

            target_producer_id = None
            if raw_p_id and raw_p_id in producer_map_id:
                target_producer_id = producer_map_id[raw_p_id]
            elif raw_p_name and raw_p_name.lower() in producer_map_name:
                target_producer_id = producer_map_name[raw_p_name.lower()]
            elif existing:
                target_producer_id = existing.get("producer_id")
            else:
                target_producer_id = fallback_producer_id

            if not target_producer_id:
                errors += 1
                error_details.append(f"Riga {r_idx} ({prod_name}): Cantina non trovata")
                continue

            v_year = None
            if vintage_idx is not None and row[vintage_idx] is not None:
                v_str = re.sub(r'\D', '', str(row[vintage_idx]))
                if v_str and len(v_str) == 4:
                    v_year = int(v_str)

            if not existing:
                q = {"producer_id": target_producer_id, "name": {"$regex": f"^{re.escape(prod_name)}$", "$options": "i"}}
                if v_year:
                    q["vintage_year"] = v_year
                existing = await db.products.find_one(q)

            is_ris = False
            if riserva_idx is not None and row[riserva_idx] is not None:
                r_val = str(row[riserva_idx]).strip().lower()
                is_ris = r_val in ["sì", "si", "true", "1", "riserva"]

            alc = None
            if alcohol_idx is not None and row[alcohol_idx] is not None:
                alc_str = str(row[alcohol_idx]).replace(',', '.')
                m_alc = re.search(r'(\d+(?:\.\d+)?)', alc_str)
                if m_alc:
                    alc = float(m_alc.group(1))

            grapes = [g.strip() for g in str(row[grapes_idx]).split(',') if g.strip()] if grapes_idx is not None and row[grapes_idx] is not None else []
            pairings = [p.strip() for p in str(row[pairings_idx]).split(',') if p.strip()] if pairings_idx is not None and row[pairings_idx] is not None else []
            photos = [ph.strip() for ph in str(row[photos_idx]).split(',') if ph.strip()] if photos_idx is not None and row[photos_idx] is not None else []

            custom_str = str(row[custom_idx]) if custom_idx is not None and row[custom_idx] is not None else ""
            custom_attrs = parse_custom_attributes_from_str(custom_str)

            # Process individual dynamic attribute columns
            for col_i, header_orig in enumerate(raw_headers_orig):
                if col_i in standard_indices:
                    continue
                if not header_orig:
                    continue
                if col_i < len(row) and row[col_i] is not None:
                    cell_val = str(row[col_i]).strip()
                    if cell_val:
                        existing_attr = next((a for a in custom_attrs if isinstance(a, dict) and a.get("name", "").lower() == header_orig.lower()), None)
                        if existing_attr:
                            existing_attr["value"] = cell_val
                        else:
                            custom_attrs.append({"name": header_orig, "value": cell_val})

            doc = {
                "name": prod_name,
                "producer_id": target_producer_id,
                "category": str(row[cat_idx]).strip().upper() if cat_idx is not None and row[cat_idx] else "VINO_ROSSO",
                "denominazione": parse_denominazione_acronym(str(row[denom_idx])) if denom_idx is not None and row[denom_idx] else "DOC",
                "vintage_year": v_year,
                "is_riserva": is_ris,
                "alcohol_degrees": alc,
                "grape_varieties": grapes,
                "food_pairings": pairings,
                "description": str(row[desc_idx]).strip() if desc_idx is not None and row[desc_idx] is not None else "",
                "tasting_notes": {
                    "visual": str(row[visual_idx]).strip() if visual_idx is not None and row[visual_idx] is not None else "",
                    "olfactory": str(row[olfactory_idx]).strip() if olfactory_idx is not None and row[olfactory_idx] is not None else "",
                    "taste": str(row[taste_idx]).strip() if taste_idx is not None and row[taste_idx] is not None else ""
                },
                "serving_temperature": str(row[temp_idx]).strip() if temp_idx is not None and row[temp_idx] is not None else "16-18°C",
                "indicative_price": str(row[price_idx]).strip() if price_idx is not None and row[price_idx] is not None else "",
                "photos": photos,
                "technical_sheet_pdf": str(row[pdf_idx]).strip() if pdf_idx is not None and row[pdf_idx] is not None else "",
                "custom_attributes": custom_attrs,
                "status": "DRAFT" if status_idx is not None and row[status_idx] and str(row[status_idx]).strip().upper() == "DRAFT" else "PUBLISHED",
                "updated_at": datetime.utcnow()
            }

            doc["slug"] = await generate_unique_product_slug(
                db=db,
                name=doc["name"],
                producer_id=target_producer_id,
                is_riserva=doc["is_riserva"],
                exclude_id=existing["_id"] if existing else None
            )

            if existing:
                await db.products.update_one({"_id": existing["_id"]}, {"$set": doc})
                updated += 1
            else:
                doc["created_at"] = datetime.utcnow()
                res = await db.products.insert_one(doc)
                created += 1

            await sync_custom_attributes_with_master(custom_attrs, db, attr_cache)
            await sync_grapes_with_master(grapes, db, grapes_cache)
            await sync_pairings_with_master(pairings, db, pairings_cache)

        except Exception as e:
            errors += 1
            error_details.append(f"Riga {r_idx}: {str(e)}")

    return {
        "message": "Importazione Excel completata",
        "created": created,
        "updated": updated,
        "errors": errors,
        "error_details": error_details[:10]
    }
